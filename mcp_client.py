"""Select and call a workflow tool exposed by mcp_server.py over stdio."""

from __future__ import annotations

import argparse
import asyncio
import csv
import json
import re
import sys
from collections.abc import Sequence
from datetime import datetime
from pathlib import Path

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parent
SERVER_PATH = ROOT / "mcp_server.py"
TEST_RESULTS_DIR = ROOT / "generate_test_results"
DEFAULT_TOOL = "create_test_environment"
TEST_SCHEMA = "mcp_test_user"
DEFAULT_ARGUMENTS = {
    "object_name": "NATL_WORKDAY_ACC_GL_FILE_LOAD_PROC",
    "connection_name": "ncrdev",
    "test_schema": TEST_SCHEMA,
}
DEFAULT_CLEANUP_ARGUMENTS = {
    "expected_schema": TEST_SCHEMA,
    "confirm": True,
}


def server_parameters() -> StdioServerParameters:
    if not SERVER_PATH.is_file():
        raise FileNotFoundError(f"Workflow MCP server not found: {SERVER_PATH}")
    return StdioServerParameters(
        command=sys.executable,
        args=[str(SERVER_PATH)],
        cwd=str(ROOT),
    )


def print_result(result) -> None:
    if getattr(result, "isError", False):
        print("Tool returned an error:", file=sys.stderr)
    for item in result.content:
        if getattr(item, "type", None) == "text":
            print(item.text)
        else:
            print(json.dumps(item.model_dump(mode="json"), indent=2))


def result_text(result) -> str:
    return "\n".join(
        item.text
        for item in result.content
        if getattr(item, "type", None) == "text"
    )


def parse_markdown_results(markdown: str) -> list[list[str]]:
    lines = [line.strip() for line in markdown.splitlines()]
    for index in range(len(lines) - 1):
        if not lines[index].startswith("|") or not lines[index + 1].startswith("|"):
            continue

        header = next(csv.reader([lines[index].strip("|")], delimiter="|"))
        separator = next(csv.reader([lines[index + 1].strip("|")], delimiter="|"))
        if not separator or not all(re.fullmatch(r"\s*:?-{3,}:?\s*", cell) for cell in separator):
            continue

        normalized_headers = [cell.strip().casefold() for cell in header]
        expected_headers = ["#", "assertion", "test input / method", "result"]
        if normalized_headers != expected_headers:
            raise ValueError(
                "Test results table has unexpected columns. Expected: "
                "#, Assertion, Test input / method, Result"
            )

        rows = [[cell.strip() for cell in header]]
        for line in lines[index + 2 :]:
            if not line.startswith("|"):
                if rows:
                    break
                continue
            cells = [cell.strip() for cell in next(csv.reader([line.strip("|")], delimiter="|"))]
            if len(cells) != len(header):
                raise ValueError("A test result row does not match the table's column count")
            rows.append(cells)
        if len(rows) == 1:
            raise ValueError("The test results table contains no assertion rows")
        return rows

    raise ValueError("The generate_test_cases response did not contain a Markdown table")


def save_test_results(markdown: str, object_name: str) -> Path:
    rows = parse_markdown_results(markdown)
    TEST_RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    safe_object_name = re.sub(r"[^A-Za-z0-9_-]+", "_", object_name).strip("_") or "test"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = TEST_RESULTS_DIR / f"{safe_object_name}_{timestamp}.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Test Results"
    for row in rows:
        worksheet.append(row)

    header_fill = PatternFill(fill_type="solid", fgColor="173B46")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    column_widths = (6, 52, 48, 72)
    for column, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[chr(64 + column)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    workbook.save(output_path)
    return output_path


async def run(args: argparse.Namespace) -> None:
    async with stdio_client(server_parameters()) as (read_stream, write_stream):
        async with ClientSession(read_stream, write_stream) as session:
            await session.initialize()
            if args.list_tools:
                result = await session.list_tools()
                for tool in result.tools:
                    print(f"{tool.name}: {tool.description or ''}".strip())
                return

            if args.tool == "cleanup":
                print(
                    f"WARNING: running destructive cleanup in {TEST_SCHEMA}; "
                    "objects listed by cleanup.sql may be permanently dropped."
                )

            requested_schema = args.arguments.get("test_schema", TEST_SCHEMA)
            if requested_schema != TEST_SCHEMA:
                raise ValueError(f"test_schema is fixed to {TEST_SCHEMA!r}")

            result = await session.call_tool(args.tool, args.arguments)
            if getattr(result, "isError", False):
                print_result(result)
                raise RuntimeError(f"Workflow tool {args.tool!r} failed")
            if args.tool == "generate_test_cases":
                output_path = save_test_results(
                    result_text(result),
                    str(args.arguments.get("object_name", "test_results")),
                )
                print(f"Test results saved to: {output_path}")
            else:
                print_result(result)


def parse_arguments(argv: Sequence[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Call a selected workflow tool from mcp_server.py."
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--list-tools",
        action="store_true",
        help="List the workflow tools exposed by mcp_server.py.",
    )
    mode.add_argument(
        "--tool",
        default=DEFAULT_TOOL,
        help=f"Exact workflow tool name to call (default: {DEFAULT_TOOL}).",
    )
    parser.add_argument(
        "--arguments",
        help="Tool arguments as a JSON object.",
    )
    args = parser.parse_args(argv)

    if args.arguments is None:
        if args.list_tools:
            args.arguments = {}
        elif args.tool == "cleanup":
            args.arguments = DEFAULT_CLEANUP_ARGUMENTS.copy()
        else:
            args.arguments = DEFAULT_ARGUMENTS.copy()

    try:
        if isinstance(args.arguments, str):
            args.arguments = json.loads(args.arguments)
    except json.JSONDecodeError as exc:
        parser.error(f"--arguments must be valid JSON: {exc}")
    if not isinstance(args.arguments, dict):
        parser.error("--arguments must contain a JSON object")
    if args.list_tools and args.arguments:
        parser.error("--arguments can only be used with --tool")
    return args


async def main(argv: Sequence[str]) -> None:
    await run(parse_arguments(argv))


if __name__ == "__main__":
    try:
        asyncio.run(main(sys.argv[1:]))
    except Exception as exc:
        print(f"Workflow MCP request failed: {exc}", file=sys.stderr)
        raise SystemExit(1) from exc